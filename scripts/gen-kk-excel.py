import json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DATA = "D:/APPMI/scripts/scan-results.json"
OUT = "D:/APPMI/LAPORAN_KK_SCAN.xlsx"

with open(DATA, encoding="utf-8") as f:
    data = json.load(f)

wb = Workbook()
wb.remove(wb.active)

HEADER_FIELDS = [
    ("Nomor KK", "nomorKK"),
    ("Nama Kepala Keluarga", "namaKepalaKeluarga"),
    ("Alamat", "alamat"),
    ("RT/RW", "rtRw"),
    ("Kode Pos", "kodePos"),
    ("Desa/Kelurahan", "kelurahan"),
    ("Kecamatan", "kecamatan"),
    ("Kabupaten/Kota", "kabupaten"),
    ("Provinsi", "provinsi"),
    ("Tanggal Dikeluarkan", "tglDikeluarkan"),
    ("KUB", "kub"),
    ("File Sumber", "file"),
    ("Confidence OCR", "confidence"),
]

MEMBER_COLS = [
    ("No", "index"),
    ("Nama Lengkap", "nama"),
    ("NIK", "nik"),
    ("Jenis Kelamin", "jenisKelamin"),
    ("Tempat Lahir", "tempatLahir"),
    ("Tanggal Lahir", "tanggalLahir"),
    ("Agama", "agama"),
    ("Pendidikan", "pendidikan"),
    ("Jenis Pekerjaan", "pekerjaan"),
    ("Golongan Darah", "golonganDarah"),
    ("Status Perkawinan", "statusPerkawinan"),
    ("Tanggal Perkawinan/Perceraian", None),
    ("Status Hubungan dalam Keluarga", "statusHubungan"),
    ("Kewarganegaraan", None),
    ("Nama Ayah", "namaAyah"),
    ("Nama Ibu", "namaIbu"),
]

def sanitize_sheet_name(name):
    base = re.sub(r"\.[^.]+$", "", name)
    base = re.sub(r"[\[\]:*?/\\]", "", base)
    base = base.strip()
    if not base:
        base = "KK"
    base = base[:31]
    return base

used_names = set()
def unique_name(name):
    if name not in used_names:
        used_names.add(name)
        return name
    i = 2
    while True:
        suffix = f" ({i})"
        cand = (name[:31-len(suffix)] + suffix)
        if cand not in used_names:
            used_names.add(cand)
            return cand
        i += 1

bold = Font(bold=True)
fill = PatternFill("solid", fgColor="D9E2F3")
wrap = Alignment(wrap_text=True, vertical="top")

for entry in data:
    raw_name = entry.get("file", "KK")
    sheet_name = unique_name(sanitize_sheet_name(raw_name))
    ws = wb.create_sheet(title=sheet_name)

    row = 1
    for label, key in HEADER_FIELDS:
        ws.cell(row=row, column=1, value=label).font = bold
        val = entry.get(key, "") if key else ""
        if val is None:
            val = ""
        ws.cell(row=row, column=2, value=val)
        row += 1

    table_start = row + 1
    for ci, (colname, _) in enumerate(MEMBER_COLS, start=1):
        c = ws.cell(row=table_start, column=ci, value=colname)
        c.font = bold
        c.fill = fill
        c.alignment = wrap
    row = table_start + 1
    for m in entry.get("members", []) or []:
        for ci, (_, key) in enumerate(MEMBER_COLS, start=1):
            if key is None:
                val = ""
            else:
                val = m.get(key, "")
            if val is None:
                val = ""
            ws.cell(row=row, column=ci, value=val).alignment = wrap
        row += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 52
    for ci in range(1, len(MEMBER_COLS)+1):
        ws.column_dimensions[get_column_letter(ci)].width = 18
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width, 28)
    ws.freeze_panes = ws.cell(row=table_start+1, column=1)

wb.save(OUT)
print("Sheets:", len(wb.sheetnames))
print("Saved:", OUT)
print("First 5:", wb.sheetnames[:5])
